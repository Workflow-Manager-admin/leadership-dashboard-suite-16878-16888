from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import BaseModel, Field
from typing import Literal
import os
from src.api.deps import get_db
from src.api.routes.auth import get_current_active_user
from src.api.security import rate_limit

from jinja2 import Template

router = APIRouter()

class ExportRequest(BaseModel):
    dashboard_id: str = Field(..., min_length=1, max_length=100, description="ID of dashboard to export")
    format: Literal["pdf", "ppt", "html"] = Field(..., description="Export format: pdf | ppt | html")

class ExportResult(BaseModel):
    url: str = Field(..., description="Download URL or confirmation")

def export_dashboard_to_html(config: dict, export_path: str):
    template = Template('''
    <html>
    <head><title>{{ title }}</title></head>
    <body>
        <h1>{{ title }}</h1>
        <pre>{{ config }}</pre>
    </body>
    </html>
    ''')
    html = template.render(title=config.get("title", "Dashboard"), config=config)
    with open(export_path, "w") as f:
        f.write(html)

def export_dashboard_to_pdf(html_path, pdf_path):
    from xhtml2pdf import pisa
    with open(html_path) as src, open(pdf_path, "w+b") as output:
        pisa.CreatePDF(src.read(), dest=output)

def export_dashboard_to_ppt(config: dict, ppt_path: str):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    shapes = slide.shapes
    shapes.title.text = config.get("title", "Dashboard Export")
    prs.save(ppt_path)

# PUBLIC_INTERFACE
@router.post(
    "/dashboard",
    summary="Export dashboard data",
    description="Export dashboard as PDF, PPT, or HTML (real file-based, robust with error handling)",
    response_model=ExportResult
)
@rate_limit(10, 60)
async def export_dashboard(
    request: Request,
    export: ExportRequest,
    db=Depends(get_db),
    current_user=Depends(get_current_active_user),
):
    dash = await db["dashboards"].find_one({"dashboard_id": export.dashboard_id})
    if not dash:
        raise HTTPException(status_code=404, detail="Dashboard not found")
    config = dash["config"]
    export_dir = "exported"
    os.makedirs(export_dir, exist_ok=True)
    fname_base = f"{export.dashboard_id}_v{dash.get('version',1)}"
    ext = export.format
    filepath = os.path.join(export_dir, f"{fname_base}.{ext}")
    url = f"/{export_dir}/{fname_base}.{ext}"
    try:
        if export.format == "html":
            export_dashboard_to_html(config, filepath)
        elif export.format == "pdf":
            # Export HTML first, then convert to PDF
            html_tmp = os.path.join(export_dir, f"{fname_base}.html")
            export_dashboard_to_html(config, html_tmp)
            export_dashboard_to_pdf(html_tmp, filepath)
        elif export.format == "ppt":
            export_dashboard_to_ppt(config, filepath)
        else:
            raise HTTPException(status_code=400, detail="Unsupported export format")
    except Exception as ex:
        raise HTTPException(status_code=500, detail=f"Export failed: {ex}")
    # Track exports in DB for audit
    await db["dashboard_exports"].insert_one({
        "dashboard_id": export.dashboard_id,
        "format": ext,
        "filename": fname_base + "." + ext,
        "url": url,
        "user": current_user["email"],
        "exported_at": os.path.getmtime(filepath) if os.path.exists(filepath) else None
    })
    return ExportResult(url=url)
