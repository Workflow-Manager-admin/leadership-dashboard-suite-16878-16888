from typing import Dict, Any
from pptx import Presentation
from fpdf import FPDF
from jinja2 import Template

# PUBLIC_INTERFACE
def dashboard_to_html(dashboard_config: Dict[str, Any]) -> str:
    """Render dashboard config to HTML string."""
    # For demo, use a very simple HTML template
    html_template = """
    <html>
      <head>
        <meta charset="utf-8">
        <title>{{ title }}</title>
        <style>
          body { font-family: Arial, sans-serif; }
          .dashboard-title { font-size: 2em; color: #0057B8; }
          .dashboard-section { margin-bottom: 2em; }
          .widget { border: 1px solid #e2e2e2; border-radius: 6px; padding: 1em; margin: 1em 0; background: #f6f7fb;}
        </style>
      </head>
      <body>
        <div class="dashboard-title">{{ title }}</div>
        <div>
          {% for widget in widgets %}
          <div class="dashboard-section">
            <div class="widget">
              <b>{{ widget.name or 'Widget' }}</b><br>
              <pre style="margin:0">{{ widget | tojson(indent=2) }}</pre>
            </div>
          </div>
          {% endfor %}
        </div>
      </body>
    </html>
    """
    # Extract title and widgets
    title = dashboard_config.get("title", "Dashboard Export")
    widgets = dashboard_config.get("widgets", [])
    # Render template
    t = Template(html_template)
    html = t.render(title=title, widgets=widgets)
    return html

# PUBLIC_INTERFACE
def generate_pdf_from_html(html: str, pdf_path: str):
    """Create a PDF file from HTML string at pdf_path."""
    # Minimalistic implementation using FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    # FPDF does not handle HTML, just as demo add as plain text
    lines = html.splitlines()
    for line in lines:
        pdf.cell(0, 10, txt=line[:100], ln=1)  # Truncate long lines, crude.
    pdf.output(pdf_path)

# PUBLIC_INTERFACE
def generate_ppt_from_dashboard(dashboard_config: Dict[str, Any], ppt_path: str):
    """Generate a PowerPoint file for the dashboard and save to ppt_path."""
    prs = Presentation()
    title = dashboard_config.get("title", "Dashboard Export")
    widgets = dashboard_config.get("widgets", [])

    # Title slide
    slide_layout = prs.slide_layouts[0]  # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Exported from SLT Dashboard"

    # One slide per widget (if exists), else info slide
    if len(widgets) == 0:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "No widgets"
        slide.placeholders[1].text = "This dashboard has no widgets configured."
    else:
        for widget in widgets:
            layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = str(widget.get("name", "Dashboard Widget"))
            body_shape = slide.placeholders[1]
            body_shape.text = str(widget)

    prs.save(ppt_path)
